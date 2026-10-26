import streamlit as st
import streamlit.components.v1 as components
import sqlite3
import fitz
import base64
import json
import os
import random
import re
import datetime
import io
import shutil
import subprocess
import tempfile
from pathlib import Path
from docx import Document
from pptx import Presentation

DB = "smart_revision_planner.db"
TODAY = datetime.date.today()


# ==================================================
# DATABASE
# ==================================================

def connect():
    conn = sqlite3.connect(DB)
    conn.row_factory = sqlite3.Row
    return conn


def init_db():
    with connect() as conn:
        conn.executescript("""
        CREATE TABLE IF NOT EXISTS subjects (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            name TEXT UNIQUE NOT NULL,
            confidence INTEGER DEFAULT 5,
            last_studied TEXT,
            created_at TEXT NOT NULL
        );

        CREATE TABLE IF NOT EXISTS pdfs (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            subject_id INTEGER NOT NULL,
            name TEXT NOT NULL,
            pdf_data BLOB NOT NULL,
            total_pages INTEGER NOT NULL,
            last_page INTEGER DEFAULT 1,
            completed INTEGER DEFAULT 0,
            uploaded_at TEXT NOT NULL,
            file_type TEXT DEFAULT 'pdf',
            UNIQUE(subject_id, name)
        );

        CREATE TABLE IF NOT EXISTS quiz_batches (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            title TEXT NOT NULL,
            scope TEXT NOT NULL,
            questions_json TEXT NOT NULL,
            score INTEGER,
            total INTEGER NOT NULL,
            answers_json TEXT,
            created_at TEXT NOT NULL,
            attempted_at TEXT
        );

        CREATE TABLE IF NOT EXISTS revision_history (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            subject_id INTEGER NOT NULL,
            activity TEXT NOT NULL,
            activity_date TEXT NOT NULL,
            details TEXT
        );

        CREATE TABLE IF NOT EXISTS notes (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            pdf_id INTEGER NOT NULL,
            page INTEGER NOT NULL,
            note TEXT NOT NULL,
            created_at TEXT NOT NULL
        );

        CREATE TABLE IF NOT EXISTS page_progress (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            material_id INTEGER NOT NULL,
            page_number INTEGER NOT NULL,
            completed INTEGER NOT NULL DEFAULT 0,
            updated_at TEXT NOT NULL,
            UNIQUE(material_id, page_number),
            FOREIGN KEY(material_id) REFERENCES pdfs(id) ON DELETE CASCADE
        );
        """)


init_db()

# Add the file_type column to existing databases without deleting saved data.
with connect() as conn:
    existing_columns = {
        row["name"] for row in conn.execute("PRAGMA table_info(pdfs)").fetchall()
    }
    if "file_type" not in existing_columns:
        conn.execute("ALTER TABLE pdfs ADD COLUMN file_type TEXT DEFAULT 'pdf'")
    conn.execute("UPDATE pdfs SET file_type = 'pdf' WHERE file_type IS NULL OR file_type = ''")
    # Preserve completion status from materials completed in older app versions.
    conn.execute("""
        INSERT OR IGNORE INTO page_progress (material_id, page_number, completed, updated_at)
        SELECT p.id, units.page_number, 1, ?
        FROM pdfs p
        JOIN (
            WITH RECURSIVE units(page_number) AS (
                SELECT 1
                UNION ALL SELECT page_number + 1 FROM units WHERE page_number < 10000
            )
            SELECT page_number FROM units
        ) units ON units.page_number <= p.total_pages
        WHERE p.completed = 1
    """, (datetime.datetime.now().isoformat(timespec="seconds"),))


# ==================================================
# GENERAL HELPERS
# ==================================================

def log_activity(subject_id, activity, details=""):
    with connect() as conn:
        conn.execute("""
            INSERT INTO revision_history
            (subject_id, activity, activity_date, details)
            VALUES (?, ?, ?, ?)
        """, (
            subject_id,
            activity,
            datetime.datetime.now().isoformat(timespec="seconds"),
            details
        ))


def get_subjects():
    with connect() as conn:
        return conn.execute("""
            SELECT * FROM subjects ORDER BY name
        """).fetchall()


def get_pdfs(subject_id=None):
    with connect() as conn:
        if subject_id is None:
            return conn.execute("""
                SELECT p.*, s.name AS subject_name
                FROM pdfs p
                JOIN subjects s ON s.id = p.subject_id
                ORDER BY s.name, p.name
            """).fetchall()

        return conn.execute("""
            SELECT * FROM pdfs
            WHERE subject_id = ?
            ORDER BY name
        """, (subject_id,)).fetchall()


def get_pdf(pdf_id):
    with connect() as conn:
        return conn.execute("""
            SELECT * FROM pdfs WHERE id = ?
        """, (pdf_id,)).fetchone()


def get_completed_pages(material_id):
    with connect() as conn:
        rows = conn.execute(
            "SELECT page_number FROM page_progress WHERE material_id = ? AND completed = 1",
            (material_id,)
        ).fetchall()
    return {int(row["page_number"]) for row in rows}


def save_page_progress(material_id, page_number, state_key):
    is_done = 1 if st.session_state.get(state_key, False) else 0
    now = datetime.datetime.now().isoformat(timespec="seconds")
    with connect() as conn:
        conn.execute(
            """INSERT INTO page_progress (material_id, page_number, completed, updated_at)
               VALUES (?, ?, ?, ?)
               ON CONFLICT(material_id, page_number) DO UPDATE SET
                 completed = excluded.completed, updated_at = excluded.updated_at""",
            (material_id, page_number, is_done, now)
        )
        total_row = conn.execute(
            "SELECT total_pages FROM pdfs WHERE id = ?", (material_id,)
        ).fetchone()
        done_row = conn.execute(
            "SELECT COUNT(*) AS done FROM page_progress WHERE material_id = ? AND completed = 1",
            (material_id,)
        ).fetchone()
        total_pages = int(total_row["total_pages"]) if total_row else 1
        all_done = int(done_row["done"]) >= total_pages
        conn.execute(
            "UPDATE pdfs SET completed = ? WHERE id = ?", (1 if all_done else 0, material_id)
        )
        material = conn.execute(
            "SELECT subject_id, name FROM pdfs WHERE id = ?", (material_id,)
        ).fetchone()
        if material:
            conn.execute(
                "UPDATE subjects SET last_studied = ? WHERE id = ?",
                (TODAY.isoformat(), material["subject_id"])
            )
    if material:
        log_activity(
            material["subject_id"],
            "Page progress updated",
            f"{material['name']} — page {page_number} {'completed' if is_done else 'reopened'}"
        )


def get_subject_stats(subject_id):
    with connect() as conn:
        rows = conn.execute(
            "SELECT id, total_pages FROM pdfs WHERE subject_id = ?", (subject_id,)
        ).fetchall()
        done_row = conn.execute(
            """SELECT COUNT(*) AS done FROM page_progress pp
               JOIN pdfs p ON p.id = pp.material_id
               WHERE p.subject_id = ? AND pp.completed = 1""",
            (subject_id,)
        ).fetchone()
    total_pages = sum(max(1, int(row["total_pages"])) for row in rows)
    done_pages = int(done_row["done"])
    percent = round(done_pages / total_pages * 100) if total_pages else 0
    return total_pages, done_pages, percent


def convert_legacy_file_to_pdf(file_bytes, filename):
    """Convert older .ppt/.doc files to PDF when LibreOffice is installed."""
    soffice = shutil.which("soffice") or shutil.which("libreoffice")
    if not soffice:
        raise ValueError(
            f"{Path(filename).suffix.upper()} is an older Office format. "
            "Install LibreOffice or save this file as .pptx/.docx and upload it again."
        )

    with tempfile.TemporaryDirectory() as tmpdir:
        source = Path(tmpdir) / Path(filename).name
        source.write_bytes(file_bytes)
        result = subprocess.run(
            [soffice, "--headless", "--convert-to", "pdf", "--outdir", tmpdir, str(source)],
            capture_output=True, text=True, timeout=60
        )
        output = Path(tmpdir) / (source.stem + ".pdf")
        if result.returncode != 0 or not output.exists():
            raise ValueError("Could not convert the older Office file. Please save it as .pptx or .docx.")
        return output.read_bytes()


def extract_document_text(file_bytes, file_type="pdf", max_chars=100000, filename=""):
    """Extract text from PDF, PowerPoint, or Word documents for quiz generation."""
    file_type = (file_type or Path(filename).suffix.lstrip(".") or "pdf").lower().lstrip(".")
    parts = []
    try:
        if file_type == "pdf":
            doc = fitz.open(stream=file_bytes, filetype="pdf")
            for i, page in enumerate(doc):
                text = page.get_text("text").strip()
                if text:
                    parts.append(f"\n--- Page {i + 1} ---\n{text}")
            doc.close()
        elif file_type == "pptx":
            presentation = Presentation(io.BytesIO(file_bytes))
            for i, slide in enumerate(presentation.slides, start=1):
                slide_parts = []
                for shape in slide.shapes:
                    if getattr(shape, "has_text_frame", False) and shape.text.strip():
                        slide_parts.append(shape.text.strip())
                    if getattr(shape, "has_table", False):
                        for row in shape.table.rows:
                            slide_parts.append(" | ".join(cell.text.strip() for cell in row.cells))
                if slide_parts:
                    parts.append(f"\n--- Slide {i} ---\n" + "\n".join(slide_parts))
        elif file_type == "docx":
            document = Document(io.BytesIO(file_bytes))
            for paragraph in document.paragraphs:
                if paragraph.text.strip():
                    parts.append(paragraph.text.strip())
            for table in document.tables:
                for row in table.rows:
                    parts.append(" | ".join(cell.text.strip() for cell in row.cells))
        elif file_type in {"ppt", "doc"}:
            converted_pdf = convert_legacy_file_to_pdf(file_bytes, filename or f"document.{file_type}")
            return extract_document_text(converted_pdf, "pdf", max_chars=max_chars)
        else:
            raise ValueError(f"Unsupported file type: .{file_type}")
        return "\n".join(parts)[:max_chars]
    except Exception as e:
        if isinstance(e, ValueError):
            raise
        raise ValueError(f"Could not extract text from {filename or file_type}: {e}") from e


def get_document_unit_label(file_type):
    return "slide" if (file_type or "pdf").lower() in {"ppt", "pptx"} else "page"


def get_reader_units(file_row):
    """Return reader units in order: PDF pages, PPT slides, or Word text pages.

    Word pagination varies by fonts and device. Without a document-layout engine,
    DOCX content is grouped into readable page-like chunks; PDFs use true pages and
    PowerPoint files use true slides.
    """
    file_type = (file_row["file_type"] or "pdf").lower()
    data = file_row["pdf_data"]
    units = []
    if file_type == "pdf":
        doc = fitz.open(stream=data, filetype="pdf")
        for index, page_obj in enumerate(doc, start=1):
            pix = page_obj.get_pixmap(matrix=fitz.Matrix(1.35, 1.35), alpha=False)
            units.append({"number": index, "kind": "image", "content": pix.tobytes("png")})
        doc.close()
    elif file_type in {"pptx", "ppt"}:
        if file_type == "ppt":
            try:
                converted = convert_legacy_file_to_pdf(data, file_row["name"])
                doc = fitz.open(stream=converted, filetype="pdf")
                for index, page_obj in enumerate(doc, start=1):
                    pix = page_obj.get_pixmap(matrix=fitz.Matrix(1.35, 1.35), alpha=False)
                    units.append({"number": index, "kind": "image", "content": pix.tobytes("png")})
                doc.close()
            except Exception as exc:
                units.append({"number": 1, "kind": "text", "content": f"Legacy .PPT preview needs LibreOffice installed. Your original file is saved and can still be downloaded.\n\nDetails: {exc}"})
        else:
            presentation = Presentation(io.BytesIO(data))
            for index, slide in enumerate(presentation.slides, start=1):
                lines = []
                for shape in slide.shapes:
                    if getattr(shape, "has_text_frame", False) and shape.text.strip():
                        lines.append(shape.text.strip())
                    if getattr(shape, "has_table", False):
                        for row in shape.table.rows:
                            lines.append(" | ".join(cell.text.strip() for cell in row.cells))
                units.append({"number": index, "kind": "text", "content": "\n\n".join(lines) or "(No selectable text on this slide.)"})
    elif file_type in {"docx", "doc"}:
        if file_type == "doc":
            try:
                converted = convert_legacy_file_to_pdf(data, file_row["name"])
                doc = fitz.open(stream=converted, filetype="pdf")
                for index, page_obj in enumerate(doc, start=1):
                    pix = page_obj.get_pixmap(matrix=fitz.Matrix(1.35, 1.35), alpha=False)
                    units.append({"number": index, "kind": "image", "content": pix.tobytes("png")})
                doc.close()
            except Exception as exc:
                units.append({"number": 1, "kind": "text", "content": f"Legacy .DOC preview needs LibreOffice installed. Your original file is saved and can still be downloaded.\n\nDetails: {exc}"})
        else:
            document = Document(io.BytesIO(data))
            blocks = [para.text.strip() for para in document.paragraphs if para.text.strip()]
            for table in document.tables:
                for row in table.rows:
                    row_text = " | ".join(cell.text.strip() for cell in row.cells).strip()
                    if row_text:
                        blocks.append(row_text)
            # Use about 8 paragraphs per readable unit when exact Word pagination is unavailable.
            chunk_size = 8
            chunks = [blocks[i:i + chunk_size] for i in range(0, len(blocks), chunk_size)] or [["(No readable text found in this document.)"]]
            for index, chunk in enumerate(chunks, start=1):
                units.append({"number": index, "kind": "text", "content": "\n\n".join(chunk)})
    return units


def get_unit_count(file_bytes, file_type, filename=""):
    """Count pages/slides/chunks using the same rules as the reader."""
    temporary_row = {"file_type": file_type, "name": filename, "pdf_data": file_bytes}
    return len(get_reader_units(temporary_row))


def display_pdf(pdf_bytes, height=750):
    """Embed a PDF in the app. Browser PDF support may vary."""
    encoded = base64.b64encode(pdf_bytes).decode("utf-8")
    components.html(
        f"""<iframe src="data:application/pdf;base64,{encoded}#toolbar=1&navpanes=1"
        width="100%" height="{height}"
        style="border:1px solid #eaded7;border-radius:10px"></iframe>""",
        height=height + 20, scrolling=True
    )


def display_document_preview(file_row):
    """Display PDF pages or readable text for PowerPoint/Word documents."""
    file_type = (file_row["file_type"] or "pdf").lower()
    data = file_row["pdf_data"]
    if file_type == "pdf":
        display_pdf(data)
        return
    if file_type in {"ppt", "doc"}:
        try:
            display_pdf(convert_legacy_file_to_pdf(data, file_row["name"]))
            return
        except Exception as exc:
            st.warning(str(exc))
    try:
        text = extract_document_text(data, file_type, max_chars=50000, filename=file_row["name"])
        if text.strip():
            st.text_area("Document text preview", text, height=500)
        else:
            st.warning("No selectable text was found in this file.")
    except Exception as exc:
        st.warning(f"Preview unavailable: {exc}")
    mime = {
        "pptx": "application/vnd.openxmlformats-officedocument.presentationml.presentation",
        "ppt": "application/vnd.ms-powerpoint",
        "docx": "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        "doc": "application/msword"
    }.get(file_type, "application/octet-stream")
    st.download_button("Download original file", data=data, file_name=file_row["name"], mime=mime)


def revision_interval(strength, days_since):
    """Weaker subjects get shorter revision intervals."""
    if strength < 35:
        return 1
    if strength < 50:
        return 2
    if strength < 65:
        return 4
    if strength < 80:
        return 7
    return 14


# ==================================================
# SUBJECT STRENGTH
# ==================================================

def subject_strength(subject_id, confidence):
    """
    Strength score:
    50% quiz accuracy + 30% confidence + 20% recency.
    Subjects without quiz results remain unassessed.
    """
    with connect() as conn:
        attempts = conn.execute("""
            SELECT score, total
            FROM quiz_batches
            WHERE score IS NOT NULL AND total > 0
        """).fetchall()

        subject = conn.execute("""
            SELECT last_studied
            FROM subjects WHERE id = ?
        """, (subject_id,)).fetchone()

        # Use quiz history from quizzes tagged to this subject.
        subject_batches = conn.execute("""
            SELECT qb.score, qb.total
            FROM quiz_batches qb
            WHERE qb.score IS NOT NULL
              AND qb.total > 0
              AND qb.scope LIKE ?
        """, (f"%subject:{subject_id}%",)).fetchall()

    if not subject_batches:
        return None

    accuracy = sum(r["score"] for r in subject_batches) / max(
        1, sum(r["total"] for r in subject_batches)
    ) * 100

    confidence_score = confidence * 10

    last_studied = subject["last_studied"] if subject else None

    if last_studied:
        try:
            last_date = datetime.date.fromisoformat(last_studied[:10])
            days = (TODAY - last_date).days
        except ValueError:
            days = 30
    else:
        days = 30

    # Recency contributes to mastery, but cannot outweigh quiz results.
    recency_score = max(0, 100 - days * 3)

    return round(
        accuracy * 0.50
        + confidence_score * 0.30
        + recency_score * 0.20
    )


def get_all_strengths():
    results = []

    for subject in get_subjects():
        strength = subject_strength(
            subject["id"], subject["confidence"]
        )
        results.append((subject, strength))

    return results


# ==================================================
# SIMPLE LOCAL QUIZ GENERATION (NO API KEY REQUIRED)
# ==================================================

def generate_questions(source_text, count=20):
    """Create basic fill-in-the-blank MCQs from extracted study text locally."""
    if not source_text or not source_text.strip():
        raise ValueError(
            "No readable text was extracted from the selected material. "
            "Try a text-based PDF/PPT/DOC file; scanned files need OCR."
        )

    stop_words = {
        "about", "above", "after", "again", "against", "being", "below",
        "between", "could", "does", "doing", "during", "each", "from",
        "have", "having", "into", "itself", "more", "most", "other",
        "over", "same", "should", "some", "such", "than", "that", "their",
        "there", "these", "they", "this", "those", "through", "under",
        "until", "very", "what", "when", "where", "which", "while",
        "with", "would", "your", "also", "because", "been", "were",
        "will", "shall", "must", "using", "used", "uses", "whose",
        "then", "them", "only", "both", "many", "much", "make", "made",
        "like", "into", "upon", "whereas", "however", "therefore", "figure",
        "table", "page", "slide", "document", "material", "chapter"
    }

    # Split text into reasonably sized sentence-like pieces, ignoring page/slide labels.
    cleaned = re.sub(r"---\s*(Page|Slide)\s+\d+\s*---", " ", source_text, flags=re.I)
    pieces = re.split(r"(?<=[.!?])\s+|[\r\n]+", cleaned)
    sentences = []
    for piece in pieces:
        piece = re.sub(r"\s+", " ", piece).strip(" \t-•")
        words = re.findall(r"\b[A-Za-z][A-Za-z'-]{3,}\b", piece)
        if len(words) >= 5 and len(piece) >= 35:
            sentences.append((piece, words))

    if not sentences:
        raise ValueError(
            "Not enough readable text to create a quiz. Select a material with more text."
        )

    # Candidate terms come from the actual material, so the quiz works offline.
    term_counts = {}
    original_terms = {}
    for sentence, words in sentences:
        for word in words:
            key = word.lower().strip("'-")
            if len(key) >= 5 and key not in stop_words and not key.isdigit():
                term_counts[key] = term_counts.get(key, 0) + 1
                original_terms.setdefault(key, word.strip("'-"))

    terms = list(term_counts)
    if len(terms) < 4:
        raise ValueError(
            "There are not enough distinct key words in this material to make multiple-choice questions."
        )

    # Prefer terms that occur more than once, but allow other terms too.
    candidates = []
    for sentence, words in sentences:
        local_keys = {w.lower().strip("'-") for w in words}
        possible = [
            w.lower().strip("'-") for w in words
            if len(w.strip("'-")) >= 5
            and w.lower().strip("'-") not in stop_words
            and w.lower().strip("'-") in term_counts
        ]
        possible = list(dict.fromkeys(possible))
        possible.sort(key=lambda k: (term_counts.get(k, 0), len(k)), reverse=True)
        for target in possible[:3]:
            candidates.append((sentence, target, local_keys))

    random.shuffle(candidates)
    questions = []
    seen = set()
    for sentence, target, local_keys in candidates:
        if len(questions) >= count:
            break
        question_text = re.sub(
            rf"\b{re.escape(target)}\b", "________", sentence, count=1, flags=re.I
        )
        if "________" not in question_text:
            continue
        signature = (question_text.lower(), target)
        if signature in seen:
            continue
        seen.add(signature)

        distractors = [t for t in terms if t != target and t not in local_keys]
        random.shuffle(distractors)
        distractors = distractors[:3]
        if len(distractors) < 3:
            distractors = [t for t in terms if t != target][:3]
        if len(distractors) < 3:
            continue

        choices = [target] + distractors
        random.shuffle(choices)
        options = {letter: original_terms.get(word, word) for letter, word in zip("ABCD", choices)}
        answer = next(letter for letter, word in zip("ABCD", choices) if word == target)
        questions.append({
            "question": f"Choose the word that best completes the sentence: {question_text}",
            "options": options,
            "answer": answer,
            "explanation": f"From the selected study material: {sentence}"
        })

    if not questions:
        raise ValueError("Could not create questions from this text. Try another material.")

    return questions


# ==================================================
# APP STYLE AND NAVIGATION
# ==================================================

st.set_page_config(
    page_title="Smart Revision Planner",
    page_icon="📚",
    layout="wide"
)

st.markdown("""
<style>
:root {
  --plum:#35334B; --plum2:#514A68; --lav:#E9E2F5; --sage:#D9E8DC;
  --rose:#F7E3E9; --blue:#E4EDF8; --paper:#FBF9FD; --white:#FFFFFF;
  --ink:#302D3D; --muted:#777386; --line:#E7E1EF;
}
html, body, [class*="css"] { font-family: Inter, ui-sans-serif, system-ui, -apple-system, BlinkMacSystemFont, "Segoe UI", sans-serif; }
.stApp { background:var(--paper); color:var(--ink); }
[data-testid="stHeader"] { background:rgba(251,249,253,.94); }
[data-testid="stMainBlockContainer"] { padding-top:2.1rem; padding-bottom:3rem; max-width:1320px; }
[data-testid="stMain"] { background:linear-gradient(180deg,#FBF9FD 0%,#F8F7FB 100%); }
/* Sidebar */
[data-testid="stSidebar"] { background:linear-gradient(180deg,#35334B 0%,#49425F 100%); border-right:1px solid #625A75; }
[data-testid="stSidebar"] [data-testid="stSidebarContent"] { padding:1.2rem .9rem; }
[data-testid="stSidebar"] h1,[data-testid="stSidebar"] h2,[data-testid="stSidebar"] h3,
[data-testid="stSidebar"] p,[data-testid="stSidebar"] label,[data-testid="stSidebar"] span { color:#F8F5FC !important; }
[data-testid="stSidebar"] h1 { font-size:1.22rem; line-height:1.4; padding-bottom:.7rem; border-bottom:1px solid #625A75; }
[data-testid="stSidebar"] [data-testid="stRadio"] > label { color:#DCD6E8 !important; font-size:.78rem; text-transform:uppercase; letter-spacing:.09em; }
[data-testid="stSidebar"] [data-testid="stRadio"] div[role="radiogroup"] { gap:.35rem; }
[data-testid="stSidebar"] [data-testid="stRadio"] div[role="radiogroup"] label { border:1px solid transparent; border-radius:12px; padding:.62rem .7rem; transition:all .18s ease; }
[data-testid="stSidebar"] [data-testid="stRadio"] div[role="radiogroup"] label:hover { background:#514A68; border-color:#716783; }
[data-testid="stSidebar"] [data-testid="stRadio"] div[role="radiogroup"] label:has(input:checked) { background:#E9E2F5; border-color:#E9E2F5; }
[data-testid="stSidebar"] [data-testid="stRadio"] div[role="radiogroup"] label:has(input:checked) * { color:#35334B !important; font-weight:700; }
[data-testid="stSidebar"] hr { border-color:#625A75; }
/* Main headings and typography */
h1 { color:var(--plum) !important; font-size:2.15rem !important; font-weight:800 !important; letter-spacing:-.045em !important; line-height:1.18 !important; margin-bottom:.35rem !important; }
h2 { color:var(--plum) !important; font-size:1.55rem !important; font-weight:760 !important; letter-spacing:-.03em !important; margin-top:1.4rem !important; }
h3 { color:var(--plum2) !important; font-size:1.12rem !important; font-weight:730 !important; }
[data-testid="stMainBlockContainer"] > div > div > div > div > p { color:var(--muted); }
p, li { line-height:1.65; }
[data-testid="stCaptionContainer"] { color:var(--muted); }
/* Give Streamlit's main content blocks a card-like surface */
[data-testid="stVerticalBlock"] > [data-testid="stVerticalBlockBorderWrapper"] { border:1px solid var(--line); border-radius:18px; background:rgba(255,255,255,.86); box-shadow:0 5px 20px rgba(53,51,75,.045); padding:1rem 1.1rem; }
[data-testid="stMetric"] { background:#fff; border:1px solid var(--line); border-radius:17px; padding:1.15rem 1.25rem; box-shadow:0 5px 18px rgba(53,51,75,.055); min-height:115px; }
[data-testid="stMetricLabel"] { color:var(--muted) !important; font-size:.88rem !important; font-weight:650 !important; }
[data-testid="stMetricValue"] { color:var(--plum) !important; font-weight:800 !important; letter-spacing:-.035em; }
[data-testid="stMetricDelta"] { font-weight:650; }
/* Buttons */
div.stButton > button, div.stDownloadButton > button, div.stFormSubmitButton > button, [data-testid="stLinkButton"] a {
  border-radius:12px; min-height:2.65rem; padding:.45rem 1rem; border:1px solid #D8CDE9;
  background:var(--lav); color:var(--plum); font-weight:720; transition:all .18s ease;
  box-shadow:0 2px 7px rgba(53,51,75,.04);
}
div.stButton > button:hover, div.stDownloadButton > button:hover, div.stFormSubmitButton > button:hover, [data-testid="stLinkButton"] a:hover {
  background:#DCD2EE; border-color:#B9A9D5; color:#302742; transform:translateY(-1px); box-shadow:0 5px 12px rgba(53,51,75,.09);
}
div.stButton > button[kind="primary"], div.stFormSubmitButton > button[kind="primary"] { background:var(--plum); border-color:var(--plum); color:white; }
div.stButton > button[kind="primary"]:hover, div.stFormSubmitButton > button[kind="primary"]:hover { background:#514A68; border-color:#514A68; color:white; }
/* Inputs */
.stTextInput input, .stTextArea textarea, .stNumberInput input, .stDateInput input,
[data-baseweb="select"] > div, [data-testid="stFileUploader"] section { border-radius:12px !important; border-color:#DED8E8 !important; background:#fff !important; }
.stTextInput input:focus, .stTextArea textarea:focus, .stNumberInput input:focus { border-color:#9B8BB9 !important; box-shadow:0 0 0 2px #E9E2F5 !important; }
[data-testid="stFileUploader"] section { background:#fff !important; border:1.5px dashed #CFC3E2 !important; padding:1.1rem; }
[data-testid="stFileUploader"] section:hover { border-color:#9986B8 !important; background:#FCFAFF !important; }
/* Tabs, progress, alerts, expanders, tables */
[data-baseweb="tab-list"] { gap:.35rem; border-bottom:1px solid var(--line); }
[data-baseweb="tab"] { border-radius:10px 10px 0 0; padding:.65rem .95rem; color:var(--muted); }
[data-baseweb="tab"][aria-selected="true"] { color:var(--plum) !important; background:var(--lav); font-weight:750; }
[data-testid="stProgressBar"] { background:#EEEAF4; border-radius:999px; }
[data-testid="stProgressBar"] > div > div { background:linear-gradient(90deg,#8F7CB6,#B6A5D4); border-radius:999px; }
div[data-testid="stAlert"] { border-radius:14px; border:1px solid var(--line); }
[data-testid="stExpander"] { border:1px solid var(--line); border-radius:14px; background:#fff; overflow:hidden; }
[data-testid="stExpander"] summary { padding:.65rem .8rem; }
[data-testid="stDataFrame"], [data-testid="stTable"] { border:1px solid var(--line); border-radius:12px; overflow:hidden; }
hr { border-color:var(--line); margin:1.5rem 0; }
/* Subjects & Study Library: tidy upload workflow and material rows */
[data-testid="stForm"] {
  background:linear-gradient(135deg,#FFFFFF 0%,#FCFAFF 100%);
  border:1px solid #E4DCEE; border-radius:18px; box-shadow:0 7px 22px rgba(53,51,75,.045);
}
[data-testid="stSelectbox"] label, [data-testid="stFileUploader"] label,
[data-testid="stTextInput"] label { color:var(--plum) !important; font-weight:700 !important; }
[data-testid="stFileUploader"] section { border-radius:16px !important; background:linear-gradient(135deg,#FFFFFF 0%,#FBF8FF 100%) !important; }
[data-testid="stExpander"] { margin:.45rem 0; border:1px solid #E4DDEE; border-radius:16px; background:#FFFFFF; box-shadow:0 4px 14px rgba(53,51,75,.035); }
[data-testid="stExpander"] summary { background:linear-gradient(90deg,#F7F3FC 0%,#FFFFFF 72%); padding:.85rem 1rem; }
[data-testid="stExpander"] summary:hover { background:#F0EAF8; }
[data-testid="stExpander"] [data-testid="stMarkdownContainer"] p { margin-bottom:.35rem; }
/* Progress analytics and completion indicators */
[data-testid="stProgressBar"] { height:12px !important; padding:2px; background:#F0ECF5 !important; border:1px solid #E5DEEF; border-radius:999px !important; overflow:hidden; }
[data-testid="stProgressBar"] > div > div { background:linear-gradient(90deg,#A99ACB 0%,#BFD8C6 100%) !important; border-radius:999px !important; }
[data-testid="stMainBlockContainer"] [data-testid="stVerticalBlockBorderWrapper"] { transition:border-color .18s ease, box-shadow .18s ease; }
[data-testid="stMainBlockContainer"] [data-testid="stVerticalBlockBorderWrapper"]:hover { border-color:#D9CDE9; box-shadow:0 7px 22px rgba(53,51,75,.065); }
/* Make status and history content easy to scan */
[data-testid="stAlert"] p { line-height:1.55; }
[data-testid="stCaptionContainer"] p { color:#777386 !important; }
/* Form and quiz elements */
[data-testid="stForm"] { border:1px solid var(--line); border-radius:16px; padding:1.15rem; background:#fff; }
[data-testid="stRadio"] > label, [data-testid="stCheckbox"] label { color:var(--ink); }
[data-testid="stRadio"] div[role="radiogroup"] { gap:.45rem; }
[data-testid="stRadio"] div[role="radiogroup"] label { border-radius:10px; padding:.4rem .55rem; }
/* Small-screen layout */
@media (max-width:768px) {
  [data-testid="stMainBlockContainer"] { padding:1.1rem .85rem 2rem; }
  h1 { font-size:1.7rem !important; }
  h2 { font-size:1.3rem !important; }
  [data-testid="stMetric"] { padding:.85rem; min-height:95px; }
  [data-testid="stVerticalBlock"] > [data-testid="stVerticalBlockBorderWrapper"] { padding:.8rem; }
}
</style>
""", unsafe_allow_html=True)

st.sidebar.title("📚 Smart Revision Planner")

page = st.sidebar.radio(
    "Navigate",
    [
        "Dashboard",
        "Subjects & Study Library",
        "Study Material Reader",
        "Quiz Centre",
        "Weak Subjects",
        "Smart Revision Planner",
        "Notes & Bookmarks",
        "Progress Analytics"
    ]
)


# ==================================================
# DASHBOARD
# ==================================================

if page == "Dashboard":
    st.title("📚 Smart Revision Planner")
    st.write("Your personal study library and revision assistant.")

    subjects = get_subjects()
    pdfs = get_pdfs()

    total_pdfs = len(pdfs)
    total_units = sum(max(1, int(p["total_pages"])) for p in pdfs)
    completed_units = sum(len(get_completed_pages(p["id"])) for p in pdfs)
    overall_percent = round(completed_units / total_units * 100) if total_units else 0

    c1, c2, c3 = st.columns(3)
    c1.metric("Subjects", len(subjects))
    c2.metric("Study Materials", total_pdfs)
    c3.metric("Reading Progress", f"{completed_units}/{total_units} pages/slides", f"{overall_percent}%")

    st.subheader("📂 Subject Progress")

    if not subjects:
        st.info("Add your first subject to get started.")

    for subject in subjects:
        total, done, percent = get_subject_stats(subject["id"])

        st.write(f"**{subject['name']}** — {done}/{total} pages/slides completed ({percent}%)")
        st.progress(percent / 100)

    st.subheader("🎯 Recommended next steps")

    strengths = get_all_strengths()
    assessed = [(s, score) for s, score in strengths if score is not None]

    if assessed:
        weakest, score = min(assessed, key=lambda item: item[1])
        st.info(
            f"Focus on **{weakest['name']}** next. "
            f"Your estimated strength is {score}/100."
        )
    elif subjects:
        st.info("Take your first quiz and rate your confidence to assess your subjects.")


# ==================================================
# SUBJECTS & PDF LIBRARY
# ==================================================

elif page == "Subjects & Study Library":
    st.title("📁 Subjects & Study Library")

    with st.form("new_subject", clear_on_submit=True):
        subject_name = st.text_input(
            "New subject name",
            placeholder="e.g. Power Electronics"
        )
        add_subject = st.form_submit_button("Create Subject")

        if add_subject:
            if not subject_name.strip():
                st.warning("Enter a subject name.")
            else:
                try:
                    with connect() as conn:
                        conn.execute("""
                            INSERT INTO subjects (name, created_at)
                            VALUES (?, ?)
                        """, (
                            subject_name.strip(),
                            datetime.datetime.now().isoformat()
                        ))
                    st.success("Subject created.")
                    st.rerun()
                except sqlite3.IntegrityError:
                    st.error("That subject already exists.")

    subjects = get_subjects()

    if not subjects:
        st.info("Create a subject before uploading study materials.")
        st.stop()

    subject_map = {s["name"]: s["id"] for s in subjects}

    st.subheader("Upload study materials into a subject folder")
    st.caption("Supported: PDF, PowerPoint (.pptx/.ppt), and Word (.docx/.doc).")

    selected_subject = st.selectbox(
        "Choose destination subject",
        list(subject_map.keys())
    )

    uploaded_files = st.file_uploader(
        "Select one or more study materials",
        type=["pdf", "pptx", "ppt", "docx", "doc"],
        accept_multiple_files=True
    )

    if st.button("Upload selected materials"):
        if not uploaded_files:
            st.warning("Select at least one file.")
        else:
            added = 0
            for uploaded in uploaded_files:
                try:
                    data = uploaded.getvalue()
                    file_type = Path(uploaded.name).suffix.lower().lstrip(".")
                    if file_type not in {"pdf", "pptx", "ppt", "docx", "doc"}:
                        raise ValueError("Unsupported file type. Upload PDF, PPT/PPTX, or DOC/DOCX.")
                    unit_count = get_unit_count(data, file_type, uploaded.name)

                    with connect() as conn:
                        conn.execute("""
                            INSERT INTO pdfs (
                                subject_id, name, pdf_data, total_pages,
                                last_page, completed, uploaded_at, file_type
                            ) VALUES (?, ?, ?, ?, 1, 0, ?, ?)
                        """, (
                            subject_map[selected_subject], uploaded.name, data,
                            max(1, unit_count), datetime.datetime.now().isoformat(), file_type
                        ))
                    added += 1
                except sqlite3.IntegrityError:
                    st.warning(f"{uploaded.name} already exists in this subject.")
                except Exception as e:
                    st.error(f"{uploaded.name}: {e}")

            if added:
                st.success(f"Uploaded {added} study material(s).")
                st.rerun()

    st.divider()

    for subject in subjects:
        with st.expander(f"📂 {subject['name']}"):
            subject_pdfs = get_pdfs(subject["id"])
            total, done, percent = get_subject_stats(subject["id"])

            st.write(f"**Reading progress:** {done}/{total} pages/slides ({percent}%)")
            st.progress(percent / 100)

            if not subject_pdfs:
                st.caption("This subject has no study materials yet.")

            for pdf in subject_pdfs:
                unit_label = get_document_unit_label(pdf["file_type"])
                completed_count = len(get_completed_pages(pdf["id"]))
                progress = round(completed_count / max(1, int(pdf["total_pages"])) * 100)
                status = "Completed" if progress == 100 else "In progress" if completed_count else "Not started"
                st.write(
                    f"📄 {pdf['name']} ({(pdf['file_type'] or 'pdf').upper()}) — {status} — "
                    f"{completed_count}/{pdf['total_pages']} {unit_label}s completed ({progress}%)"
                )


# ==================================================
# CONTINUOUS STUDY READER
# ==================================================

elif page == "Study Material Reader":
    st.title("📖 Reading Mode")
    st.caption("Choose a subject and material once, then scroll through it naturally. Mark each page or slide beside its content; progress saves automatically.")

    subjects = get_subjects()
    if not subjects:
        st.info("Create a subject and upload study materials first.")
        st.stop()

    subject_map = {subject["name"]: subject["id"] for subject in subjects}
    chosen_subject = st.selectbox("Choose subject", list(subject_map.keys()), key="reader_subject")
    subject_id = subject_map[chosen_subject]
    materials = get_pdfs(subject_id)
    if not materials:
        st.info("This subject has no study materials yet. Upload one in Subjects & Study Library.")
        st.stop()

    # Include the ID in the label so duplicate filenames remain distinguishable.
    material_labels = {f"{m['name']} · {m['id']}": m["id"] for m in materials}
    chosen_label = st.selectbox("Open study material", list(material_labels.keys()), key=f"reader_material_{subject_id}")
    material = get_pdf(material_labels[chosen_label])

    try:
        units = get_reader_units(material)
    except Exception as exc:
        st.error(f"Could not open this material: {exc}")
        st.stop()

    # Keep the saved total aligned with the actual reader units (e.g. Word text chunks).
    actual_total = max(1, len(units))
    if int(material["total_pages"]) != actual_total:
        with connect() as conn:
            conn.execute("UPDATE pdfs SET total_pages = ? WHERE id = ?", (actual_total, material["id"]))
        material = get_pdf(material["id"])

    completed_units = get_completed_pages(material["id"])
    done_count = sum(1 for unit in units if unit["number"] in completed_units)
    percent = round(done_count / actual_total * 100)
    metric1, metric2 = st.columns([1, 2])
    metric1.metric("Pages/slides completed", f"{done_count}/{actual_total}")
    metric2.metric("Reading progress", f"{percent}%")
    st.progress(percent / 100)
    st.divider()
    st.subheader(material["name"])

    # One continuous Streamlit page: each page/slide and its checkbox are rendered together.
    unit_label = get_document_unit_label(material["file_type"])
    for unit in units:
        number = unit["number"]
        with st.container(border=True):
            content_col, check_col = st.columns([5, 1], vertical_alignment="top")
            with content_col:
                st.markdown(f"**{unit_label.title()} {number}**")
                if unit["kind"] == "image":
                    st.image(unit["content"], use_container_width=True)
                else:
                    st.markdown(unit["content"].replace("\n", "  \n"))
            with check_col:
                state_key = f"page_done_{material['id']}_{number}"
                if state_key not in st.session_state:
                    st.session_state[state_key] = number in completed_units
                st.checkbox(
                    "Completed",
                    key=state_key,
                    on_change=save_page_progress,
                    args=(material["id"], number, state_key)
                )

    # Stats refresh on each checkbox interaction; all values are persisted in SQLite.
    latest_done = len(get_completed_pages(material["id"]))
    latest_percent = round(latest_done / actual_total * 100)
    st.divider()
    st.write(f"**Material progress:** {latest_done}/{actual_total} {unit_label}s completed ({latest_percent}%).")
    st.progress(latest_percent / 100)
    if latest_done == actual_total:
        st.success("🎉 All pages/slides in this material are completed!")
    if (material["file_type"] or "").lower() == "docx":
        st.caption("Note: Word page breaks depend on formatting. Without a Word layout engine, DOCX is shown as readable page-like chunks; completion is saved per chunk.")
    elif (material["file_type"] or "").lower() == "pptx":
        st.caption("PowerPoint is displayed slide-by-slide as extracted text. Original slide layouts and embedded media may not be reproduced in this reader.")


# ==================================================
# QUIZ CENTRE
# ==================================================

elif page == "Quiz Centre":
    st.title("🧠 Quiz Centre")

    pdfs = get_pdfs()
    subjects = get_subjects()

    if not pdfs:
        st.info("Upload study materials before generating a quiz.")
        st.stop()

    quiz_mode = st.radio(
        "Quiz scope",
        [
            "One Material",
            "Selected Materials",
            "Entire Subject",
            "Full Syllabus"
        ],
        horizontal=True
    )

    source_pdfs = []
    scope_tag = ""
    title = ""

    if quiz_mode == "One Material":
        pdf_map = {
            f"{p['subject_name']} / {p['name']}": p
            for p in pdfs
        }
        chosen = st.selectbox("Choose material", list(pdf_map.keys()))
        source_pdfs = [pdf_map[chosen]]
        scope_tag = f"pdf:{source_pdfs[0]['id']}"
        title = source_pdfs[0]["name"]

    elif quiz_mode == "Selected Materials":
        pdf_map = {
            f"{p['subject_name']} / {p['name']}": p
            for p in pdfs
        }
        selected = st.multiselect(
            "Choose materials to combine",
            list(pdf_map.keys())
        )
        source_pdfs = [pdf_map[name] for name in selected]
        scope_tag = "selected:" + ",".join(
            str(p["id"]) for p in source_pdfs
        )
        title = "Combined Study Material Quiz"

    elif quiz_mode == "Entire Subject":
        subject_map = {s["name"]: s["id"] for s in subjects}
        subject_name = st.selectbox(
            "Choose subject", list(subject_map.keys())
        )
        subject_id = subject_map[subject_name]
        source_pdfs = [
            p for p in pdfs if p["subject_id"] == subject_id
        ]
        scope_tag = f"subject:{subject_id}"
        title = f"{subject_name} Quiz"

    else:
        source_pdfs = pdfs
        scope_tag = "full-syllabus"
        title = "Full Syllabus Quiz"

    st.write(f"**Study materials included:** {len(source_pdfs)}")

    if source_pdfs:
        with st.expander("View included materials"):
            for p in source_pdfs:
                st.write(f"- {p['subject_name']} / {p['name']}")

    if st.button("✨ Generate Up to 20 Questions", type="primary"):
        if not source_pdfs:
            st.warning("Select at least one study material.")
        else:
            combined_text = ""

            try:
                with st.spinner("Extracting study material..."):
                    for p in source_pdfs:
                        text = extract_document_text(
                            p["pdf_data"], p["file_type"] or "pdf",
                            max_chars=30000, filename=p["name"]
                        )
                        combined_text += f"\n\nDOCUMENT: {p['name']}\n{text}"

                combined_text = combined_text[:100000]

                with st.spinner("Creating questions from your study material..."):
                    questions = generate_questions(combined_text, 20)

                with connect() as conn:
                    cur = conn.execute("""
                        INSERT INTO quiz_batches (
                            title, scope, questions_json,
                            total, created_at
                        )
                        VALUES (?, ?, ?, ?, ?)
                    """, (
                        title,
                        scope_tag,
                        json.dumps(questions),
                        len(questions),
                        datetime.datetime.now().isoformat()
                    ))
                    batch_id = cur.lastrowid

                st.session_state["active_quiz_id"] = batch_id
                st.success(f"Created {len(questions)} questions. Your previous batches are preserved.")
                st.rerun()

            except Exception as e:
                st.error(f"Quiz generation failed: {e}")
                st.info("Check that the selected material contains readable text. This quiz generator works locally and does not need an API key.")

    st.divider()
    st.subheader("🗂️ Previous Quiz Batches")

    with connect() as conn:
        batches = conn.execute("""
            SELECT * FROM quiz_batches
            ORDER BY id DESC
        """).fetchall()

    if not batches:
        st.info("Your generated quizzes will appear here.")
    else:
        batch_map = {
            f"#{b['id']} — {b['title']} — "
            f"{b['created_at'][:16]}": b["id"]
            for b in batches
        }

        selected_batch_label = st.selectbox(
            "Open a quiz batch",
            list(batch_map.keys())
        )
        batch_id = batch_map[selected_batch_label]

        with connect() as conn:
            batch = conn.execute("""
                SELECT * FROM quiz_batches WHERE id = ?
            """, (batch_id,)).fetchone()

        questions = json.loads(batch["questions_json"])

        if batch["score"] is not None:
            st.success(
                f"Previous score: {batch['score']}/{batch['total']} "
                f"({batch['score'] / batch['total'] * 100:.0f}%)"
            )

            if batch["answers_json"]:
                saved_answers = json.loads(batch["answers_json"])

                with st.expander("Review answers"):
                    for i, q in enumerate(questions):
                        user_answer = saved_answers.get(str(i), "")
                        correct = q["answer"]

                        st.write(f"**Q{i + 1}. {q['question']}**")
                        st.write(f"Your answer: {user_answer or 'Skipped'}")
                        st.write(f"Correct answer: {correct}")
                        st.write(q.get("explanation", ""))

        else:
            with st.form(f"quiz_form_{batch_id}"):
                answers = {}

                for i, q in enumerate(questions):
                    st.markdown(f"**Q{i + 1}. {q['question']}**")

                    choices = list(q["options"].items())

                    answers[str(i)] = st.radio(
                        f"Select answer for question {i + 1}",
                        choices,
                        format_func=lambda item: (
                            f"{item[0]}. {item[1]}"
                        ),
                        index=None,
                        key=f"batch_{batch_id}_q_{i}",
                        label_visibility="collapsed"
                    )

                submit = st.form_submit_button("Submit Quiz")

                if submit:
                    score = 0
                    answer_letters = {}

                    for i, q in enumerate(questions):
                        chosen_answer = answers[str(i)]
                        letter = chosen_answer[0] if chosen_answer else ""
                        answer_letters[str(i)] = letter

                        if letter == q["answer"]:
                            score += 1

                    with connect() as conn:
                        conn.execute("""
                            UPDATE quiz_batches
                            SET score = ?,
                                answers_json = ?,
                                attempted_at = ?
                            WHERE id = ?
                        """, (
                            score,
                            json.dumps(answer_letters),
                            datetime.datetime.now().isoformat(),
                            batch_id
                        ))

                    # Update subject study dates for quizzes tagged to a subject.
                    if batch["scope"].startswith("subject:"):
                        try:
                            subject_id = int(
                                batch["scope"].split(":")[1]
                            )
                            with connect() as conn:
                                conn.execute("""
                                    UPDATE subjects SET last_studied = ?
                                    WHERE id = ?
                                """, (TODAY.isoformat(), subject_id))
                        except ValueError:
                            pass

                    st.success(
                        f"Your score: {score}/{len(questions)} "
                        f"({score / len(questions) * 100:.0f}%)"
                    )
                    st.rerun()


# ==================================================
# WEAK SUBJECTS
# ==================================================

elif page == "Weak Subjects":
    st.title("📉 Subjects: Weakest to Strongest")

    subjects = get_subjects()

    if not subjects:
        st.info("Create subjects first.")
        st.stop()

    st.subheader("Rate your confidence")

    for subject in subjects:
        with st.container(border=True):
            st.write(f"**{subject['name']}**")

            confidence = st.slider(
                "Confidence (1 = very weak, 10 = very strong)",
                1, 10,
                value=int(subject["confidence"]),
                key=f"subject_conf_{subject['id']}"
            )

            if st.button(
                "Save confidence",
                key=f"save_conf_{subject['id']}"
            ):
                with connect() as conn:
                    conn.execute("""
                        UPDATE subjects SET confidence = ?
                        WHERE id = ?
                    """, (confidence, subject["id"]))

                st.success("Confidence saved.")
                st.rerun()

    st.divider()
    st.subheader("Strength Ranking")

    strengths = get_all_strengths()
    assessed = [(s, score) for s, score in strengths if score is not None]
    unassessed = [(s, score) for s, score in strengths if score is None]

    assessed.sort(key=lambda item: item[1])

    for subject, score in assessed:
        st.write(f"**{subject['name']}** — Strength: {score}/100")
        st.progress(score / 100)

    for subject, _ in unassessed:
        st.info(
            f"{subject['name']} — Not assessed yet. "
            "Complete a quiz to estimate strength."
        )


# ==================================================
# SMART REVISION PLANNER
# ==================================================

elif page == "Smart Revision Planner":
    st.title("🗓️ Smart Revision Planner")

    study_hours = st.slider(
        "Available study time today (hours)",
        min_value=1,
        max_value=12,
        value=3
    )

    subjects = get_subjects()

    if not subjects:
        st.info("Add subjects first.")
        st.stop()

    plan = []

    for subject in subjects:
        strength = subject_strength(
            subject["id"], subject["confidence"]
        )

        last_studied = subject["last_studied"]

        if last_studied:
            try:
                last_date = datetime.date.fromisoformat(last_studied[:10])
                days = max(0, (TODAY - last_date).days)
            except ValueError:
                days = 30
        else:
            days = 30

        if strength is None:
            priority = 50 + min(days, 30)
            reason = "Not assessed or no quiz score yet"
        else:
            priority = (100 - strength) * 0.7 + min(days, 30) * 1.0
            reason = f"Strength {strength}/100; last studied {days} days ago"

        plan.append({
            "name": subject["name"],
            "id": subject["id"],
            "strength": strength,
            "days": days,
            "priority": priority,
            "reason": reason
        })

    plan.sort(key=lambda item: item["priority"], reverse=True)

    st.caption(
        "The planner prioritizes weak subjects and those you have not "
        "studied recently. Adjust the daily hours to fit your schedule."
    )

    minutes_available = study_hours * 60
    minutes_per_subject = min(
        60,
        max(20, minutes_available // max(1, min(len(plan), 4)))
    )

    for i, item in enumerate(plan):
        if i >= 4:
            break

        with st.container(border=True):
            st.subheader(f"{i + 1}. {item['name']}")
            st.write(item["reason"])
            st.write(f"Suggested session: {minutes_per_subject} minutes")

            if item["strength"] is not None:
                st.progress(item["strength"] / 100)

            if st.button(
                f"Mark {item['name']} as studied today",
                key=f"studied_{item['id']}"
            ):
                with connect() as conn:
                    conn.execute("""
                        UPDATE subjects SET last_studied = ?
                        WHERE id = ?
                    """, (TODAY.isoformat(), item["id"]))

                log_activity(
                    item["id"],
                    "Revision session",
                    "Marked studied from Smart Revision Planner"
                )

                st.success("Study session recorded.")
                st.rerun()


# ==================================================
# NOTES & BOOKMARKS
# ==================================================

elif page == "Notes & Bookmarks":
    st.title("📝 Notes & Bookmarks")

    pdfs = get_pdfs()

    if not pdfs:
        st.info("Upload a study material first.")
        st.stop()

    pdf_map = {
        f"{p['subject_name']} / {p['name']}": p
        for p in pdfs
    }

    chosen = st.selectbox("Choose material", list(pdf_map.keys()))
    pdf = pdf_map[chosen]

    with st.form("save_note", clear_on_submit=True):
        page_number = st.number_input(
            "Page number",
            min_value=1,
            max_value=pdf["total_pages"],
            value=min(pdf["last_page"], pdf["total_pages"])
        )

        note_text = st.text_area(
            "Note, formula, doubt, or bookmark description"
        )

        save_note = st.form_submit_button("Save Note")

        if save_note:
            if not note_text.strip():
                st.warning("Enter a note.")
            else:
                with connect() as conn:
                    conn.execute("""
                        INSERT INTO notes (
                            pdf_id, page, note, created_at
                        )
                        VALUES (?, ?, ?, ?)
                    """, (
                        pdf["id"],
                        int(page_number),
                        note_text.strip(),
                        datetime.datetime.now().isoformat()
                    ))

                st.success("Note saved.")
                st.rerun()

    st.divider()
    st.subheader("Saved Notes")

    with connect() as conn:
        notes = conn.execute("""
            SELECT n.*, p.name AS pdf_name
            FROM notes n
            JOIN pdfs p ON p.id = n.pdf_id
            WHERE n.pdf_id = ?
            ORDER BY n.id DESC
        """, (pdf["id"],)).fetchall()

    for note in notes:
        with st.expander(f"Page {note['page']} — {note['created_at'][:10]}"):
            st.write(note["note"])

            if st.button("Delete note", key=f"delete_note_{note['id']}"):
                with connect() as conn:
                    conn.execute(
                        "DELETE FROM notes WHERE id = ?",
                        (note["id"],)
                    )
                st.rerun()


# ==================================================
# PROGRESS ANALYTICS
# ==================================================

elif page == "Progress Analytics":
    st.title("📊 Progress Analytics")

    subjects = get_subjects()
    pdfs = get_pdfs()

    if not subjects:
        st.info("Add subjects to see analytics.")
        st.stop()

    st.subheader("Study Material Completion")

    for subject in subjects:
        total, done, percent = get_subject_stats(subject["id"])
        with st.container(border=True):
            left, right = st.columns([4, 1])
            with left:
                st.markdown(f"**📘 {subject['name']}**")
                st.caption(f"{done} of {total} pages/slides completed")
            with right:
                st.markdown(f"<div style='text-align:right;color:#35334B;font-size:1.25rem;font-weight:800;padding-top:.15rem'>{percent}%</div>", unsafe_allow_html=True)
            st.progress(percent / 100)

    st.divider()
    st.subheader("Quiz Performance")

    with connect() as conn:
        batches = conn.execute("""
            SELECT title, score, total, attempted_at
            FROM quiz_batches
            WHERE score IS NOT NULL
            ORDER BY id DESC
        """).fetchall()

    if not batches:
        st.info("Complete a quiz to see performance analytics.")
    else:
        for batch in batches:
            percentage = round(batch["score"] / batch["total"] * 100)
            with st.container(border=True):
                top, score = st.columns([4, 1])
                with top:
                    st.markdown(f"**📝 {batch['title']}**")
                    st.caption(batch['attempted_at'] or 'Date unavailable')
                with score:
                    st.markdown(f"<div style='text-align:right;color:#35334B;font-weight:800;font-size:1.1rem'>{batch['score']}/{batch['total']}</div>", unsafe_allow_html=True)
                st.progress(percentage / 100)
                st.caption(f"Quiz score: {percentage}%")

    st.divider()
    st.subheader("Recent Study Activity")

    with connect() as conn:
        history = conn.execute("""
            SELECT h.activity_date, s.name, h.activity, h.details
            FROM revision_history h
            JOIN subjects s ON s.id = h.subject_id
            ORDER BY h.id DESC
            LIMIT 30
        """).fetchall()

    if not history:
        st.info("Your reading and revision activity will appear here.")
    else:
        for row in history:
            with st.container(border=True):
                st.markdown(f"**🗓️ {row['activity_date'][:16]}** · {row['name']}")
                st.write(f"{row['activity']} {'— ' + row['details'] if row['details'] else ''}")
